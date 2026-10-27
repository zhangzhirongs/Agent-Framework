import io
import zipfile

import pytest

from app.config import Settings
from app.skills import Skills
from tests.conftest import browser, login


async def test_auth_csrf_and_user_file_isolation(platform):
    app, _ = platform
    app.state.store.create_user("member", "member-test-password")
    async with browser(app) as admin, browser(app) as member, browser(app) as anonymous:
        assert (await anonymous.get("/api/threads")).status_code == 401
        await login(admin)
        await login(member, "member", "member-test-password")
        assert (await admin.post("/api/threads", headers={"X-CSRF-Token": ""})).status_code == 403
        assert (
            await admin.get("/api/threads", headers={"Origin": "https://foreign.example"})
        ).status_code == 403
        thread = (await admin.post("/api/threads")).json()["id"]
        file = (
            await admin.post(
                f"/api/threads/{thread}/files", files={"file": ("note.txt", b"private")}
            )
        ).json()
        assert (await member.get(f"/api/threads/{thread}")).status_code == 404
        assert (await member.get(f"/api/files/{file['id']}")).status_code == 404
        assert (await member.get("/api/connectors/config")).status_code == 403
        assert (
            await member.put("/api/connectors/config", json={"mcpServers": {}})
        ).status_code == 403
        assert (
            await member.post("/api/skills/import", files={"files": ("SKILL.md", b"bad")})
        ).status_code == 403
        assert (await admin.get(f"/api/files/{file['id']}/text")).json() == {"text": "private"}


async def test_model_config_is_admin_only_and_persisted(platform):
    app, _ = platform
    app.state.runtime.settings.model = "env-default"
    app.state.runtime.settings.models = "env-option,"
    app.state.runtime.settings.vision_model = "env-vision"
    app.state.store.create_user("model-member", "member-test-password")
    async with browser(app) as admin, browser(app) as member:
        await login(admin)
        await login(member, "model-member", "member-test-password")
        config = {"models": ["chat-model", "vision-model"], "vision_model": "vision-model"}
        assert (await member.get("/api/models/config")).status_code == 403
        assert (await member.put("/api/models/config", json=config)).status_code == 403
        assert (await admin.put("/api/models/config", json=config)).status_code == 200
        caps = (await member.get("/api/capabilities")).json()
        assert "chat-model" in caps["models"]
        assert "vision-model" in caps["models"]
        assert "" not in caps["models"]
        configured = (await admin.get("/api/models/config")).json()
        assert configured["managed_models"] == ["chat-model", "vision-model"]
        assert configured["environment_models"] == [
            "env-default",
            "env-option",
            "env-vision",
        ]
        assert set(app.state.runtime.available_models) == set(caps["models"])


async def test_selected_model_is_forwarded_to_runtime(platform, monkeypatch):
    app, _ = platform
    app.state.store.set_config("models", "alternate-model")
    app.state.runtime.available_models.append("alternate-model")
    selected = []
    monkeypatch.setattr(
        app.state.runtime,
        "start",
        lambda run, prompt, files, skills, geoflow_context, model_name: selected.append(model_name),
    )
    async with browser(app) as client:
        await login(client)
        thread = (await client.post("/api/threads")).json()["id"]
        response = await client.post(
            f"/api/threads/{thread}/messages",
            json={"content": "Use another model", "model": "alternate-model"},
        )
        assert response.status_code == 202
        assert selected == ["alternate-model"]


async def test_password_change_revokes_sessions_and_demo_is_not_reset(platform, settings):
    app, _ = platform
    async with browser(app) as client, browser(app) as other:
        await login(client)
        await login(other)
        response = await client.post(
            "/api/auth/password",
            json={"current_password": "123", "new_password": "updated-test-password"},
        )
        assert response.status_code == 200
        assert (await other.get("/api/auth/me")).status_code == 401
        app.state.store.bootstrap_demo()
        assert (
            await client.post("/api/auth/login", json={"username": "admin", "password": "123"})
        ).status_code == 401
        await login(client, password="updated-test-password")


async def test_html_preview_is_sandboxed_and_file_types_are_checked(platform):
    app, _ = platform
    async with browser(app) as client:
        await login(client)
        thread = (await client.post("/api/threads")).json()["id"]
        invalid = await client.post(
            f"/api/threads/{thread}/files", files={"file": ("app.exe", b"invalid")}
        )
        assert invalid.status_code == 422
        invalid_image = await client.post(
            f"/api/threads/{thread}/files", files={"file": ("image.png", b"not an image")}
        )
        assert invalid_image.status_code == 422
        html = (
            await client.post(
                f"/api/threads/{thread}/files",
                files={"file": ("index.html", b"<h1>Hi</h1><script>alert(1)</script>")},
            )
        ).json()
        preview = await client.get(f"/api/files/{html['id']}?preview=true")
        assert "sandbox" in preview.headers["content-security-policy"]
        assert preview.headers["x-content-type-options"] == "nosniff"


async def test_docx_rich_preview_preserves_headings_runs_and_tables(platform):
    import io

    from docx import Document

    from tests.conftest import browser, login

    app, _ = platform
    document = Document()
    document.add_heading("Quarterly report", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Important").bold = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    output = io.BytesIO()
    document.save(output)

    async with browser(app) as client:
        await login(client)
        thread = await client.post("/api/threads")
        uploaded = await client.post(
            f"/api/threads/{thread.json()['id']}/files",
            files={
                "file": (
                    "report.docx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        preview = await client.get(f"/api/files/{uploaded.json()['id']}/preview")

    assert preview.status_code == 200
    assert "<h1" in preview.json()["html"]
    assert "<strong>Important</strong>" in preview.json()["html"]
    assert "<table>" in preview.json()["html"]


async def test_csv_rich_preview_handles_quotes_newlines_and_user_isolation(platform):
    app, _ = platform
    app.state.store.create_user("reader", "reader-test-password")
    async with browser(app) as client, browser(app) as other:
        await login(client)
        await login(other, "reader", "reader-test-password")
        thread = (await client.post("/api/threads")).json()["id"]
        uploaded = await client.post(
            f"/api/threads/{thread}/files",
            files={
                "file": (
                    "table.csv",
                    '\ufeff名称,说明\r\n"A,B","第一行\n第二行"\r\n<script>,=1+1'.encode(),
                )
            },
        )
        file_id = uploaded.json()["id"]
        response = await client.get(f"/api/files/{file_id}/preview")
        assert response.status_code == 200
        html = response.json()["html"]
        assert "<thead><tr><th>名称</th><th>说明</th></tr></thead>" in html
        assert "<td>A,B</td>" in html
        assert "第一行<br>第二行" in html
        assert "&lt;script&gt;" in html
        assert "<td>=1+1</td>" in html
        assert (await other.get(f"/api/files/{file_id}/preview")).status_code == 404


def test_csv_preview_limits_and_empty_files():
    from app.files import extract_csv_preview

    assert "此文件内容为空" in extract_csv_preview({"content": b""})
    rows = "\n".join(str(index) for index in range(1002))
    html = extract_csv_preview({"content": rows.encode()})
    assert "1000 行、200 列" in html
    assert "<td>999</td>" in html
    assert "<td>1000</td>" not in html
    wide = extract_csv_preview({"content": ",".join(str(i) for i in range(201)).encode()})
    assert "<th>199</th>" in wide
    assert "<th>200</th>" not in wide


async def test_markdown_extension_upload_and_preview_type_check(platform):
    app, _ = platform
    async with browser(app) as client:
        await login(client)
        thread = (await client.post("/api/threads")).json()["id"]
        uploaded = await client.post(
            f"/api/threads/{thread}/files",
            files={"file": ("article.markdown", "# 标题\n\n**正文**".encode())},
        )
        assert uploaded.status_code == 201
        file_id = uploaded.json()["id"]
        assert (await client.get(f"/api/files/{file_id}/text")).json()["text"].startswith("# 标题")
        assert (await client.get(f"/api/files/{file_id}/preview")).status_code == 415


def test_production_forbids_demo_password_and_requires_secure_cookie():
    with pytest.raises(ValueError):
        Settings(_env_file=None, deployment_mode="production")
    Settings(_env_file=None, deployment_mode="production", demo_account=False, secure_cookie=True)


async def test_production_rejects_an_existing_default_password(settings):
    from app.api import create_app
    from app.store import Store

    store = Store(settings.data_dir)
    store.bootstrap_demo()
    settings.deployment_mode = "production"
    settings.demo_account = False
    settings.secure_cookie = True
    with pytest.raises(ValueError, match="admin"):
        create_app(settings)


@pytest.mark.parametrize(
    "config",
    [
        {"url": 123},
        {"command": "example", "args": [123]},
        {"url": "https://example.test/mcp", "enabled": "false"},
        {"command": "example", "allowed_tools": "all"},
        {"command": "example", "env": {"KEY": "inline-secret"}},
    ],
)
def test_mcp_rejects_invalid_configuration_without_saving(tmp_path, config):
    from app.connectors import Connectors

    registry = Connectors(tmp_path / "mcp.json")
    with pytest.raises(ValueError):
        registry.save({"example": config})
    assert not registry.path.exists()


def test_skills_standard_format_resources_and_path_traversal(tmp_path):
    skills = Skills(tmp_path / "skills")
    content = (
        b"---\nname: sample\ndescription: sample description\nmetadata:\n"
        b"  author: Example Author\n---\nInstructions"
    )
    skills.install([("sample/SKILL.md", content), ("sample/references/note.md", b"resource")])
    assert skills.catalog()[0]["author"] == "Example Author"
    assert skills.resource("sample", "references/note.md") == "resource"
    assert skills.detail("sample")["resources"] == ["SKILL.md", "references/note.md"]
    with pytest.raises((KeyError, ValueError)):
        skills.resource("sample", "../../outside.txt")
    with pytest.raises(ValueError):
        skills.install([("../SKILL.md", content)])
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as zipped:
        zipped.writestr("../SKILL.md", content)
    with pytest.raises(ValueError):
        skills.install_zip(stream.getvalue())
    skills.remove("sample")
    assert skills.catalog() == []


async def test_unconfigured_model_fails_before_accepting_message(settings):
    from app.api import create_app

    app = create_app(settings)
    async with browser(app) as client:
        await login(client)
        thread = (await client.post("/api/threads")).json()["id"]
        response = await client.post(f"/api/threads/{thread}/messages", json={"content": "hello"})
        assert response.status_code == 503
        assert not app.state.store.query("SELECT * FROM runs")
    await app.state.runtime.close()
