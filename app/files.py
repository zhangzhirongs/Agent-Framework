import html
import io
import mimetypes
import re
import time
from pathlib import PurePosixPath

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
from pypdf import PdfReader

from app.store import new_id

ALLOWED = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
}


def save_file(store, thread_id, name, content):
    name = PurePosixPath(name.replace("\\", "/")).name
    if not name or len(name) > 180 or re.search(r"[\x00-\x1f]", name):
        raise ValueError("文件名称无效")
    suffix = PurePosixPath(name).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("暂不支持此文件类型")
    media_type = mimetypes.guess_type(name)[0] or "application/octet-stream"
    if media_type.startswith("image/"):
        try:
            with Image.open(io.BytesIO(content)) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("图片像素数量超过限制")
                image.verify()
                expected = {
                    ".png": "PNG",
                    ".jpg": "JPEG",
                    ".jpeg": "JPEG",
                    ".webp": "WEBP",
                    ".gif": "GIF",
                }
                if image.format != expected[suffix]:
                    raise ValueError("图片内容与后缀不一致")
        except (OSError, SyntaxError, Image.DecompressionBombError):
            raise ValueError("图片内容无效或过大") from None
    identity = new_id()
    with store.connect() as db:
        db.execute(
            "INSERT INTO files VALUES(?,?,?,?,?,?)",
            (identity, thread_id, name, media_type, content, time.time()),
        )
    return {"id": identity, "name": name, "media_type": media_type}


def extract_text(item):
    content = item["content"]
    name = item["name"].lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages[:80])[:100000]
    if name.endswith(".docx"):
        # Enforce a decompression budget before the document parser opens the archive.
        import zipfile

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("文档解压大小超出限制")
        document = Document(io.BytesIO(content))
        paragraphs = [p.text for p in document.paragraphs]
        paragraphs += [
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs)[:100000]
    if item["media_type"].startswith("image/"):
        return "此文件为图片；启用视觉模型后可随消息分析。"
    return content.decode("utf-8", errors="replace")[:100000]


def extract_docx_preview(item):
    """Render DOCX paragraphs and tables as escaped, self-contained preview HTML."""
    import zipfile

    content = item["content"]
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 40 * 1024 * 1024:
            raise ValueError("鏂囨。瑙ｅ帇澶у皬瓒呭嚭闄愬埗")
    document = Document(io.BytesIO(content))
    rendered = []
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            value = "".join(
                _render_docx_run(run)
                for run in paragraph.runs
            ) or "&nbsp;"
            style_name = paragraph.style.name if paragraph.style else ""
            heading = re.match(r"Heading ([1-6])", style_name)
            tag = f"h{heading.group(1)}" if heading else "p"
            alignment = {
                WD_ALIGN_PARAGRAPH.CENTER: "center",
                WD_ALIGN_PARAGRAPH.RIGHT: "right",
                WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
            }.get(paragraph.alignment, "left")
            rendered.append(f'<{tag} style="text-align:{alignment}">{value}</{tag}>')
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            rows = []
            for row_index, row in enumerate(table.rows):
                cell_tag = "th" if row_index == 0 else "td"
                cells = "".join(
                    f"<{cell_tag}>{html.escape(cell.text).replace(chr(10), '<br>')}</{cell_tag}>"
                    for cell in row.cells
                )
                rows.append(f"<tr>{cells}</tr>")
            rendered.append(f"<table><tbody>{''.join(rows)}</tbody></table>")
    body_html = "".join(rendered)
    return (
        '<!doctype html><html><head><meta charset="utf-8"><style>'
        "body{margin:0;padding:36px;background:#e9edf0;color:#253238;"
        "font:15px/1.75 Arial,sans-serif}"
        '.page{max-width:760px;min-height:900px;margin:auto;padding:64px 72px;background:white;'
        'box-shadow:0 4px 24px #17252b20;overflow-wrap:anywhere}h1,h2,h3,h4,h5,h6{line-height:1.35}'
        'table{width:100%;border-collapse:collapse;margin:18px 0}td,th{border:1px solid #aebbc1;'
        'padding:7px 10px;text-align:left}th{background:#edf3f5}p{margin:0 0 12px}'
        '@media(max-width:640px){body{padding:12px}.page{padding:28px 22px;min-height:0}}'
        '</style></head><body><main class="page">'
        + body_html
        + "</main></body></html>"
    )


def _render_docx_run(run):
    value = html.escape(run.text).replace("\n", "<br>").replace("\t", "&emsp;")
    if run.bold:
        value = f"<strong>{value}</strong>"
    if run.italic:
        value = f"<em>{value}</em>"
    if run.underline:
        value = f"<u>{value}</u>"
    if run.font.color and run.font.color.rgb:
        value = f'<span style="color:#{run.font.color.rgb}">{value}</span>'
    if run.font.highlight_color:
        value = f"<mark>{value}</mark>"
    return value
